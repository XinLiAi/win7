#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
设备密码管理系统 v3.0（多用户权限版）
====================================
角色与权限：
  管理员(admin)   ：管理用户（创建/编辑/删除/重置密码，三种角色账号均只能由管理员创建）
                   + 查看【用户登录日志】；不参与设备增删改查
  操作员(operator)：设备增删改查 + Excel 导入导出；不能创建账号、不能看日志
  审计员(auditor) ：只读查看【操作日志】（设备增删改查等）+ 导出日志；无任何写操作

安全特性：
  - 首次运行强制创建管理员账号，之后必须用户名+密码登录
  - 登录密码 PBKDF2-HMAC-SHA256 加盐哈希
  - 设备密码加密存储（CTR 流加密，密钥文件权限 600），列表默认脱敏
  - 全程审计：登录成功/失败/退出 → 登录日志；设备与用户操作 → 操作日志
  - 兼容 CentOS 7（Python 3.6+，仅标准库；openpyxl 可选）
"""

import os
import sys
import re
import time
import csv
import base64
import hashlib
import hmac
import sqlite3
import tkinter as tk
import tkinter.font as tkfont
from tkinter import ttk, messagebox, filedialog, simpledialog

# 可选依赖：openpyxl（CentOS 7 可 pip3 install openpyxl）
try:
    from openpyxl import Workbook, load_workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# ---------- 全局配置 ----------
APP_NAME = "设备密码管理系统"
APP_VERSION = "v3.1 UI优化版"
APP_DIR = os.path.expanduser("~/.device_manager")
DB_FILE = os.path.join(APP_DIR, "devices.db")
KEY_FILE = os.path.join(APP_DIR, "secret.key")
ENC_PREFIX = "enc:"
PBKDF2_ITERATIONS = 120000
MIN_PASSWORD_LEN = 8
MAX_LOGIN_ATTEMPTS = 3
DEVICE_TYPES = ["服务器", "网络设备", "安防设备"]

# 角色定义
ROLE_ADMIN = "admin"
ROLE_OPERATOR = "operator"
ROLE_AUDITOR = "auditor"
ROLE_CN = {ROLE_ADMIN: "管理员", ROLE_OPERATOR: "操作员", ROLE_AUDITOR: "审计员"}
ROLE_COLORS = {ROLE_ADMIN: "#D97706", ROLE_OPERATOR: "#1E88E5", ROLE_AUDITOR: "#16A34A"}
ALL_ROLES = [ROLE_ADMIN, ROLE_OPERATOR, ROLE_AUDITOR]

os.makedirs(APP_DIR, exist_ok=True)


# ============================================================
# 一、安全模块（标准库实现）
# ============================================================
def hash_password(password):
    salt = os.urandom(16)
    dk = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ITERATIONS)
    return "pbkdf2$%d$%s$%s" % (
        PBKDF2_ITERATIONS,
        base64.b64encode(salt).decode("ascii"),
        base64.b64encode(dk).decode("ascii"),
    )


def verify_password(password, stored):
    try:
        parts = stored.split("$")
        if len(parts) != 4 or parts[0] != "pbkdf2":
            return False
        salt = base64.b64decode(parts[2])
        expected = base64.b64decode(parts[3])
        dk = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, int(parts[1]))
        return hmac.compare_digest(dk, expected)
    except Exception:
        return False


def get_or_create_key():
    if not os.path.exists(KEY_FILE):
        key = os.urandom(32)
        fd = os.open(KEY_FILE, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        with os.fdopen(fd, "wb") as f:
            f.write(key)
    with open(KEY_FILE, "rb") as f:
        return f.read()


def _keystream(key, nonce, length):
    out = b""
    counter = 0
    while len(out) < length:
        out += hmac.new(key, nonce + counter.to_bytes(8, "big"), hashlib.sha256).digest()
        counter += 1
    return out[:length]


def encrypt_text(key, plaintext):
    if not plaintext:
        return ""
    data = plaintext.encode("utf-8")
    nonce = os.urandom(16)
    ks = _keystream(key, nonce, len(data))
    cipher = bytes(a ^ b for a, b in zip(data, ks))
    return ENC_PREFIX + base64.b64encode(nonce + cipher).decode("ascii")


def decrypt_text(key, token):
    if not token:
        return ""
    if not token.startswith(ENC_PREFIX):
        return token
    try:
        raw = base64.b64decode(token[len(ENC_PREFIX):].encode("ascii"))
        nonce, cipher = raw[:16], raw[16:]
        ks = _keystream(key, nonce, len(cipher))
        return bytes(a ^ b for a, b in zip(cipher, ks)).decode("utf-8")
    except Exception:
        return ""


def now_str():
    return time.strftime("%Y-%m-%d %H:%M:%S")


# ============================================================
# 二、数据库初始化（设备 / 用户 / 两类日志）
# ============================================================
def init_db():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    # 设备表
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS devices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            room TEXT NOT NULL,
            rack TEXT NOT NULL,
            name TEXT NOT NULL,
            type TEXT NOT NULL,
            ip TEXT DEFAULT '',
            port TEXT DEFAULT '',
            account TEXT DEFAULT '',
            password TEXT DEFAULT '',
            notes TEXT DEFAULT ''
        )
    ''')
    cursor.execute("PRAGMA table_info(devices)")
    columns = [col[1] for col in cursor.fetchall()]
    if "ip" not in columns:
        cursor.execute("ALTER TABLE devices ADD COLUMN ip TEXT DEFAULT ''")
    if "port" not in columns:
        cursor.execute("ALTER TABLE devices ADD COLUMN port TEXT DEFAULT ''")
    if "account" not in columns:
        if "username" in columns:
            cursor.execute("ALTER TABLE devices ADD COLUMN account TEXT DEFAULT ''")
            cursor.execute("UPDATE devices SET account = username")
        else:
            cursor.execute("ALTER TABLE devices ADD COLUMN account TEXT DEFAULT ''")
    # 用户表
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL,
            enabled INTEGER DEFAULT 1,
            created_at TEXT,
            created_by TEXT,
            last_login TEXT
        )
    ''')
    # 登录日志（管理员可见）
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS login_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ts TEXT NOT NULL,
            username TEXT NOT NULL,
            role TEXT NOT NULL,
            result TEXT NOT NULL,
            detail TEXT DEFAULT ''
        )
    ''')
    # 操作日志（审计员可见）
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS operation_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ts TEXT NOT NULL,
            username TEXT NOT NULL,
            role TEXT NOT NULL,
            action TEXT NOT NULL,
            target TEXT DEFAULT '',
            detail TEXT DEFAULT ''
        )
    ''')
    conn.commit()
    conn.close()


def migrate_plaintext_passwords(key):
    """旧版明文设备密码迁移为加密存储"""
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("SELECT id, password FROM devices WHERE password IS NOT NULL AND password != ''")
    rows = cursor.fetchall()
    updated = 0
    for device_id, pwd in rows:
        if pwd.startswith(ENC_PREFIX):
            continue
        cursor.execute("UPDATE devices SET password=? WHERE id=?",
                       (encrypt_text(key, pwd), device_id))
        updated += 1
    conn.commit()
    conn.close()
    return updated


# ============================================================
# 三、用户管理
# ============================================================
def count_users():
    conn = sqlite3.connect(DB_FILE)
    n = conn.execute("SELECT COUNT(*) FROM users").fetchone()[0]
    conn.close()
    return n


def get_user_by_username(username):
    conn = sqlite3.connect(DB_FILE)
    row = conn.execute("SELECT id, username, password_hash, role, enabled, created_at, created_by, last_login "
                       "FROM users WHERE username=?", (username,)).fetchone()
    conn.close()
    if not row:
        return None
    return {"id": row[0], "username": row[1], "password_hash": row[2], "role": row[3],
            "enabled": row[4], "created_at": row[5], "created_by": row[6], "last_login": row[7]}


def username_exists(username):
    return get_user_by_username(username) is not None


def create_user(username, password, role, created_by):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("INSERT INTO users (username, password_hash, role, enabled, created_at, created_by) "
                 "VALUES (?, ?, ?, 1, ?, ?)",
                 (username, hash_password(password), role, now_str(), created_by))
    conn.commit()
    conn.close()


def list_users():
    conn = sqlite3.connect(DB_FILE)
    rows = conn.execute("SELECT id, username, role, enabled, created_at, created_by, last_login "
                        "FROM users ORDER BY id").fetchall()
    conn.close()
    return rows


def update_user_role(user_id, role):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("UPDATE users SET role=? WHERE id=?", (role, user_id))
    conn.commit()
    conn.close()


def set_user_enabled(user_id, enabled):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("UPDATE users SET enabled=? WHERE id=?", (1 if enabled else 0, user_id))
    conn.commit()
    conn.close()


def set_user_password(user_id, password):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("UPDATE users SET password_hash=? WHERE id=?", (hash_password(password), user_id))
    conn.commit()
    conn.close()


def delete_user(user_id):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("DELETE FROM users WHERE id=?", (user_id,))
    conn.commit()
    conn.close()


def touch_last_login(user_id):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("UPDATE users SET last_login=? WHERE id=?", (now_str(), user_id))
    conn.commit()
    conn.close()


def enabled_admin_count():
    conn = sqlite3.connect(DB_FILE)
    n = conn.execute("SELECT COUNT(*) FROM users WHERE role=? AND enabled=1", (ROLE_ADMIN,)).fetchone()[0]
    conn.close()
    return n


# ============================================================
# 四、审计日志
# ============================================================
def login_log(username, role, result, detail=""):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("INSERT INTO login_logs (ts, username, role, result, detail) VALUES (?, ?, ?, ?, ?)",
                 (now_str(), username, role, result, detail))
    conn.commit()
    conn.close()


def op_log(username, role, action, target="", detail=""):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("INSERT INTO operation_logs (ts, username, role, action, target, detail) VALUES (?, ?, ?, ?, ?, ?)",
                 (now_str(), username, role, action, target, detail))
    conn.commit()
    conn.close()


def list_login_logs(search=""):
    conn = sqlite3.connect(DB_FILE)
    if search:
        like = "%" + search + "%"
        rows = conn.execute("SELECT ts, username, role, result, detail FROM login_logs "
                            "WHERE ts LIKE ? OR username LIKE ? OR role LIKE ? OR result LIKE ? OR detail LIKE ? "
                            "ORDER BY id DESC", (like, like, like, like, like)).fetchall()
    else:
        rows = conn.execute("SELECT ts, username, role, result, detail FROM login_logs "
                            "ORDER BY id DESC").fetchall()
    conn.close()
    return rows


def list_operation_logs(search=""):
    conn = sqlite3.connect(DB_FILE)
    if search:
        like = "%" + search + "%"
        rows = conn.execute("SELECT ts, username, role, action, target, detail FROM operation_logs "
                            "WHERE ts LIKE ? OR username LIKE ? OR role LIKE ? OR action LIKE ? "
                            "OR target LIKE ? OR detail LIKE ? ORDER BY id DESC",
                            (like, like, like, like, like, like)).fetchall()
    else:
        rows = conn.execute("SELECT ts, username, role, action, target, detail FROM operation_logs "
                            "ORDER BY id DESC").fetchall()
    conn.close()
    return rows


def clear_login_logs():
    conn = sqlite3.connect(DB_FILE)
    conn.execute("DELETE FROM login_logs")
    conn.commit()
    conn.close()


def export_operation_logs_to_csv(path):
    rows = list_operation_logs()
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["时间", "用户", "角色", "操作", "对象", "详情"])
        for r in rows:
            writer.writerow(list(r))
    return len(rows)


def export_operation_logs_to_excel(path):
    rows = list_operation_logs()
    wb = Workbook()
    ws = wb.active
    ws.title = "操作日志"
    ws.append(["时间", "用户", "角色", "操作", "对象", "详情"])
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return len(rows)


# ============================================================
# 五、首次建号 / 登录
# ============================================================
def _password_strength_ok(pwd):
    return (len(pwd) >= MIN_PASSWORD_LEN
            and re.search(r"[A-Za-z]", pwd)
            and re.search(r"\d", pwd))


def _validate_password_dialog(dialog, pwd, confirm):
    if not pwd or not confirm:
        messagebox.showerror("错误", "密码不能为空！", parent=dialog)
        return False
    if not _password_strength_ok(pwd):
        messagebox.showerror("密码强度不足",
                             "密码至少 %d 位，且需同时包含字母和数字！" % MIN_PASSWORD_LEN,
                             parent=dialog)
        return False
    if pwd != confirm:
        messagebox.showerror("错误", "两次输入的密码不一致！", parent=dialog)
        return False
    return True


def _make_modal(parent, title, width, height):
    dlg = tk.Toplevel(parent)
    dlg.title(title)
    dlg.configure(bg="#F1F5F9")
    # 注意：父窗口隐藏(如登录前 withdraw)时不能设 transient，
    # 否则 update 后对话框会被连带隐藏。可见时才设为父窗口的瞬态窗口。
    if parent.winfo_viewable():
        dlg.transient(parent)
    dlg.grab_set()
    dlg.resizable(False, False)
    dlg.update_idletasks()
    try:
        if parent.winfo_viewable():
            x = parent.winfo_rootx() + max((parent.winfo_width() - width) // 2, 0)
            y = parent.winfo_rooty() + max((parent.winfo_height() - height) // 2, 0)
        else:
            x = max((dlg.winfo_screenwidth() - width) // 2, 0)
            y = max((dlg.winfo_screenheight() - height) // 2, 0)
    except Exception:
        x = max((dlg.winfo_screenwidth() - width) // 2, 0)
        y = max((dlg.winfo_screenheight() - height) // 2, 0)
    dlg.geometry("%dx%d+%d+%d" % (width, height, x, y))
    return dlg


def _fit_center(dlg, min_w=400, min_h=None):
    """让 Toplevel 对话框按内容自动定宽高并居中（略偏上，视觉更平衡）。

    关键：不再用固定像素高度，而是取内容实际需要的尺寸；
    同时用 minsize 锁定最小尺寸，防止窗口管理器把窗口压缩、导致底部按钮被裁掉。
    """
    dlg.update_idletasks()
    w = max(dlg.winfo_reqwidth(), min_w)
    h = dlg.winfo_reqheight()
    if min_h:
        h = max(h, min_h)
    sw, sh = dlg.winfo_screenwidth(), dlg.winfo_screenheight()
    x = (sw - w) // 2
    y = max((sh - h) // 3, 0)   # 略偏上，不贴着顶部也不居中偏低
    dlg.geometry("%dx%d+%d+%d" % (w, h, x, y))
    dlg.minsize(w, h)           # 锁定最小尺寸，禁止被压缩
    dlg.update_idletasks()


def initial_admin_dialog(parent):
    """首次运行：强制创建管理员账号"""
    dlg = tk.Toplevel(parent)
    dlg.title("首次使用 - 创建管理员账号")
    dlg.configure(bg="#1F2937")
    if parent.winfo_viewable():
        dlg.transient(parent)
    dlg.grab_set()
    dlg.resizable(False, False)

    tk.Label(dlg, text="首次使用，请创建管理员账号", bg="#1F2937", fg="#FFFFFF",
             font=("TkDefaultFont", 14, "bold")).pack(pady=(22, 4))
    tk.Label(dlg, text="只有管理员能创建用户，请妥善保管此账号密码", bg="#1F2937", fg="#93C5FD",
             font=("TkDefaultFont", 9)).pack(pady=(0, 10))

    body = tk.Frame(dlg, bg="#FFFFFF", highlightbackground="#E2E8F0", highlightthickness=1)
    body.pack(fill=tk.BOTH, expand=True, padx=18, pady=(4, 16))

    tk.Label(body, text="管理员用户名 *", bg="#FFFFFF", fg="#1E293B",
             font=("TkDefaultFont", 10)).pack(anchor=tk.W, padx=20, pady=(16, 4))
    user_var = tk.StringVar(value="admin")
    tk.Entry(body, textvariable=user_var, width=30, font=("TkDefaultFont", 11),
             relief="solid", bd=1).pack(fill=tk.X, padx=20, pady=(0, 8), ipady=4)

    tk.Label(body, text="登录密码 *（至少%d位，含字母和数字）" % MIN_PASSWORD_LEN,
             bg="#FFFFFF", fg="#1E293B", font=("TkDefaultFont", 10)).pack(anchor=tk.W, padx=20, pady=(4, 2))
    pwd_var = tk.StringVar()
    tk.Entry(body, textvariable=pwd_var, show="*", width=30, font=("TkDefaultFont", 11),
             relief="solid", bd=1).pack(fill=tk.X, padx=20, pady=(0, 8), ipady=4)

    tk.Label(body, text="确认密码 *", bg="#FFFFFF", fg="#1E293B",
             font=("TkDefaultFont", 10)).pack(anchor=tk.W, padx=20, pady=(4, 2))
    cfm_var = tk.StringVar()
    tk.Entry(body, textvariable=cfm_var, show="*", width=30, font=("TkDefaultFont", 11),
             relief="solid", bd=1).pack(fill=tk.X, padx=20, pady=(0, 8), ipady=4)

    result = {"ok": False}

    def on_create():
        uname = user_var.get().strip()
        if not uname:
            messagebox.showerror("错误", "管理员用户名不能为空！", parent=dlg)
            return
        if username_exists(uname):
            messagebox.showerror("错误", "用户名“%s”已存在！" % uname, parent=dlg)
            return
        if not _validate_password_dialog(dlg, pwd_var.get(), cfm_var.get()):
            return
        create_user(uname, pwd_var.get(), ROLE_ADMIN, "系统")
        login_log(uname, ROLE_CN[ROLE_ADMIN], "创建初始管理员账号", "首个管理员账号")
        result["ok"] = True
        dlg.destroy()

    btns = tk.Frame(body, bg="#FFFFFF")
    btns.pack(fill=tk.X, padx=20, pady=(14, 16))
    tk.Button(btns, text="创 建", command=on_create, bg="#1E88E5", fg="#FFFFFF",
              activebackground="#1565C0", activeforeground="#FFFFFF", relief="flat",
              padx=18, pady=6, font=("TkDefaultFont", 10, "bold")).pack(side=tk.RIGHT, padx=(8, 0))
    tk.Button(btns, text="退 出", command=dlg.destroy, bg="#FFFFFF", fg="#1E293B",
              activebackground="#E2E8F0", relief="solid", bd=1, padx=18, pady=6,
              font=("TkDefaultFont", 10)).pack(side=tk.RIGHT)

    _fit_center(dlg)
    dlg.wait_window()
    return result["ok"]


def login_dialog(parent):
    """用户名 + 密码登录"""
    dlg = tk.Toplevel(parent)
    dlg.title("登录 - %s" % APP_NAME)
    dlg.configure(bg="#1F2937")
    if parent.winfo_viewable():
        dlg.transient(parent)
    dlg.grab_set()
    dlg.resizable(False, False)

    # 顶部标题带（深色、高度均衡）
    tk.Label(dlg, text="设备密码管理系统", bg="#1F2937", fg="#FFFFFF",
             font=("TkDefaultFont", 16, "bold")).pack(pady=(22, 2))
    tk.Label(dlg, text=APP_VERSION + " · 用户登录", bg="#1F2937", fg="#93C5FD",
             font=("TkDefaultFont", 9)).pack(pady=(0, 8))

    # 表单卡片
    body = tk.Frame(dlg, bg="#FFFFFF", highlightbackground="#E2E8F0", highlightthickness=1)
    body.pack(fill=tk.BOTH, expand=True, padx=18, pady=(4, 16))

    tk.Label(body, text="用户名", bg="#FFFFFF", fg="#1E293B",
             font=("TkDefaultFont", 10)).pack(anchor=tk.W, padx=20, pady=(16, 4))
    user_var = tk.StringVar()
    user_entry = tk.Entry(body, textvariable=user_var, width=30, font=("TkDefaultFont", 11),
                          relief="solid", bd=1)
    user_entry.pack(fill=tk.X, padx=20, pady=(0, 8), ipady=4)
    user_entry.focus_set()

    tk.Label(body, text="密码", bg="#FFFFFF", fg="#1E293B",
             font=("TkDefaultFont", 10)).pack(anchor=tk.W, padx=20, pady=(4, 2))
    pwd_var = tk.StringVar()
    pwd_entry = tk.Entry(body, textvariable=pwd_var, show="*", width=30, font=("TkDefaultFont", 11),
                         relief="solid", bd=1)
    pwd_entry.pack(fill=tk.X, padx=20, pady=(0, 8), ipady=4)

    hint = tk.Label(body, text="", bg="#FFFFFF", fg="#DC2626", font=("TkDefaultFont", 9))
    hint.pack(anchor=tk.W, padx=20, pady=(2, 0))

    result = {"user": None}
    attempts = [0]

    def do_login(event=None):
        uname = user_var.get().strip()
        pwd = pwd_var.get()
        if not uname or not pwd:
            hint.config(text="请输入用户名和密码")
            return
        user = get_user_by_username(uname)
        if not user:
            attempts[0] += 1
            login_log(uname, "未知", "登录失败", "用户不存在")
            hint.config(text="用户名或密码错误（剩余 %d 次）" % (MAX_LOGIN_ATTEMPTS - attempts[0]))
        elif not user["enabled"]:
            login_log(uname, ROLE_CN.get(user["role"], user["role"]), "登录失败", "账号已禁用")
            hint.config(text="该账号已被禁用，请联系管理员")
        elif not verify_password(pwd, user["password_hash"]):
            attempts[0] += 1
            login_log(uname, ROLE_CN.get(user["role"], user["role"]), "登录失败", "密码错误")
            hint.config(text="用户名或密码错误（剩余 %d 次）" % (MAX_LOGIN_ATTEMPTS - attempts[0]))
        else:
            touch_last_login(user["id"])
            login_log(uname, ROLE_CN.get(user["role"], user["role"]), "登录成功", "")
            result["user"] = {"id": user["id"], "username": uname, "role": user["role"]}
            dlg.destroy()
            return
        if attempts[0] >= MAX_LOGIN_ATTEMPTS:
            messagebox.showerror("登录失败", "连续 %d 次登录失败，程序退出！" % MAX_LOGIN_ATTEMPTS, parent=dlg)
            dlg.destroy()

    btns = tk.Frame(body, bg="#FFFFFF")
    btns.pack(fill=tk.X, padx=20, pady=(14, 16))
    tk.Button(btns, text="登 录", command=do_login, bg="#1E88E5", fg="#FFFFFF",
              activebackground="#1565C0", activeforeground="#FFFFFF", relief="flat",
              padx=20, pady=6, font=("TkDefaultFont", 10, "bold")).pack(side=tk.RIGHT, padx=(8, 0))
    tk.Button(btns, text="退 出", command=dlg.destroy, bg="#FFFFFF", fg="#1E293B",
              activebackground="#E2E8F0", relief="solid", bd=1, padx=20, pady=6,
              font=("TkDefaultFont", 10)).pack(side=tk.RIGHT)

    _fit_center(dlg)
    dlg.bind("<Return>", do_login)
    dlg.wait_window()
    return result["user"]


def clear_children(widget):
    for c in widget.winfo_children():
        c.destroy()


# ============================================================
# 六、设备 CRUD
# ============================================================
def fetch_all_devices():
    conn = sqlite3.connect(DB_FILE)
    rows = conn.execute("SELECT id, room, rack, name, type, ip, port, account, password, notes "
                        "FROM devices ORDER BY id").fetchall()
    conn.close()
    return rows


def add_device(room, rack, name, type_, ip="", port="", account="", password="", notes=""):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("INSERT INTO devices (room, rack, name, type, ip, port, account, password, notes) "
                 "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                 (room, rack, name, type_, ip, port, account, password, notes))
    conn.commit()
    conn.close()


def update_device(device_id, room, rack, name, type_, ip="", port="", account="", password="", notes=""):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("UPDATE devices SET room=?, rack=?, name=?, type=?, ip=?, port=?, account=?, password=?, notes=? "
                 "WHERE id=?", (room, rack, name, type_, ip, port, account, password, notes, device_id))
    conn.commit()
    conn.close()


def delete_device(device_id):
    conn = sqlite3.connect(DB_FILE)
    conn.execute("DELETE FROM devices WHERE id=?", (device_id,))
    conn.commit()
    conn.close()


def get_device_by_id(device_id):
    conn = sqlite3.connect(DB_FILE)
    row = conn.execute("SELECT room, rack, name, type, ip, port, account, password, notes "
                       "FROM devices WHERE id=?", (device_id,)).fetchone()
    conn.close()
    return row


# ============================================================
# 七、主界面（角色化）
# ============================================================
class DeviceManagerApp:
    C_HEADER = "#1F2937"
    C_BG = "#F1F5F9"
    C_CARD = "#FFFFFF"
    C_ZEBRA = "#F8FAFC"
    C_TEXT = "#1E293B"
    C_MUTED = "#94A3B8"
    C_BORDER = "#E2E8F0"
    C_PRIMARY = "#1E88E5"
    C_PRIMARY_DARK = "#1565C0"
    C_SELECT = "#BFDBFE"
    C_DANGER = "#DC2626"
    C_DANGER_DARK = "#B91C1C"
    C_SUCCESS = "#16A34A"

    def __init__(self, root, user):
        self.root = root
        self.user = user                       # {'id','username','role'}
        self.role = user["role"]
        self.role_cn = ROLE_CN[self.role]
        self.key = get_or_create_key()
        self.show_pwd = tk.BooleanVar(value=False)
        self._logout_requested = False
        self.font_family = self._pick_font_family()
        self.total_count = 0
        self.filtered_count = 0

        self.root.title("%s %s - 当前用户：%s（%s）" % (APP_NAME, APP_VERSION, user["username"], self.role_cn))
        self.root.geometry("1120x680")
        self.root.minsize(920, 560)
        self.root.configure(bg=self.C_BG)

        self._setup_style()
        self._build_header()
        self._build_nav()
        self.content = tk.Frame(self.root, bg=self.C_BG)
        self.content.pack(fill=tk.BOTH, expand=True, padx=12, pady=4)
        self._build_statusbar()
        self.show_view(self._default_view())

    # ---------- 工具 ----------
    def _pick_font_family(self):
        candidates = ["Noto Sans CJK SC", "WenQuanYi Micro Hei", "WenQuanYi Zen Hei",
                      "Microsoft YaHei", "PingFang SC", "SimHei", "DejaVu Sans"]
        try:
            families = set(tkfont.families(self.root))
            for f in candidates:
                if f in families:
                    return f
        except Exception:
            pass
        return "TkDefaultFont"

    def _setup_style(self):
        self.style = ttk.Style(self.root)
        try:
            self.style.theme_use("clam")
        except tk.TclError:
            pass
        self.style.configure(".", background=self.C_BG, foreground=self.C_TEXT,
                             font=(self.font_family, 10))
        self.style.configure("Tool.TButton", background=self.C_PRIMARY, foreground="#FFFFFF",
                             borderwidth=0, padding=(14, 7), font=(self.font_family, 10, "bold"))
        self.style.map("Tool.TButton",
                       background=[("active", self.C_PRIMARY_DARK), ("pressed", self.C_PRIMARY_DARK)],
                       foreground=[("active", "#FFFFFF")])
        self.style.configure("Danger.TButton", background=self.C_DANGER, foreground="#FFFFFF",
                             borderwidth=0, padding=(14, 7), font=(self.font_family, 10, "bold"))
        self.style.map("Danger.TButton",
                       background=[("active", self.C_DANGER_DARK), ("pressed", self.C_DANGER_DARK)],
                       foreground=[("active", "#FFFFFF")])
        self.style.configure("Ghost.TButton", background="#FFFFFF", foreground=self.C_PRIMARY,
                             borderwidth=1, relief="solid", padding=(12, 6), font=(self.font_family, 10))
        self.style.map("Ghost.TButton", background=[("active", "#EFF6FF")],
                       foreground=[("active", self.C_PRIMARY_DARK)])
        self.style.configure("Nav.TButton", background="#FFFFFF", foreground="#475569",
                             borderwidth=1, relief="solid", padding=(16, 7), font=(self.font_family, 10))
        self.style.map("Nav.TButton", background=[("active", "#E2E8F0")],
                       foreground=[("active", "#1E293B")])
        self.style.configure("NavSel.TButton", background=self.C_PRIMARY, foreground="#FFFFFF",
                             borderwidth=0, padding=(16, 7), font=(self.font_family, 10, "bold"))
        self.style.map("NavSel.TButton", background=[("active", self.C_PRIMARY_DARK)],
                       foreground=[("active", "#FFFFFF")])
        self.style.configure("Search.TEntry", fieldbackground="#FFFFFF",
                             bordercolor=self.C_BORDER, padding=(8, 6), insertcolor=self.C_TEXT)
        self.style.configure("TCheckbutton", background=self.C_BG, foreground=self.C_TEXT)
        self.style.map("TCheckbutton", background=[("active", self.C_BG)])
        self.style.configure("Treeview", background=self.C_CARD, fieldbackground=self.C_CARD,
                             foreground=self.C_TEXT, borderwidth=0, relief="flat",
                             rowheight=36, font=(self.font_family, 10))
        self.style.map("Treeview", background=[("selected", self.C_SELECT)],
                       foreground=[("selected", self.C_TEXT)])
        self.style.configure("Treeview.Heading", background=self.C_PRIMARY, foreground="#FFFFFF",
                             borderwidth=0, relief="flat", padding=(6, 7),
                             font=(self.font_family, 10, "bold"))
        self.style.map("Treeview.Heading", background=[("active", self.C_PRIMARY_DARK)])
        self.style.configure("Vertical.TScrollbar", background="#CBD5E1",
                             troughcolor=self.C_BG, borderwidth=0, arrowsize=12)

    # ---------- Header ----------
    def _build_header(self):
        header = tk.Frame(self.root, bg=self.C_HEADER, height=56)
        header.pack(side=tk.TOP, fill=tk.X)
        header.pack_propagate(False)
        tk.Label(header, text=APP_NAME, bg=self.C_HEADER, fg="#FFFFFF",
                 font=(self.font_family, 16, "bold")).pack(side=tk.LEFT, padx=16, pady=12)
        tk.Label(header, text=APP_VERSION, bg=self.C_HEADER, fg="#93C5FD",
                 font=(self.font_family, 10)).pack(side=tk.LEFT, pady=15)

        tk.Button(header, text="退出登录", command=self._logout,
                  bg="#374151", fg="#FFFFFF", activebackground="#4B5563",
                  activeforeground="#FFFFFF", relief="flat", padx=12, pady=3,
                  font=(self.font_family, 9)).pack(side=tk.RIGHT, padx=14, pady=12)

        week_map = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
        today = time.strftime("%Y-%m-%d") + " " + week_map[int(time.strftime("%w")) - 1]
        tk.Label(header, text=today, bg=self.C_HEADER, fg="#CBD5E1",
                 font=(self.font_family, 10)).pack(side=tk.RIGHT, padx=10, pady=15)

        tk.Label(header, text="当前用户：%s" % self.user["username"], bg=self.C_HEADER, fg="#E2E8F0",
                 font=(self.font_family, 10)).pack(side=tk.RIGHT, padx=6, pady=15)
        tk.Label(header, text=self.role_cn, bg=ROLE_COLORS[self.role], fg="#FFFFFF",
                 font=(self.font_family, 9, "bold"), padx=8, pady=2).pack(side=tk.RIGHT, pady=13)

    def _logout(self):
        if messagebox.askyesno("退出登录", "确定要退出当前账号吗？", parent=self.root):
            login_log(self.user["username"], self.role_cn, "退出登录", "")
            self._logout_requested = True
            self.root.quit()

    # ---------- 导航（按角色） ----------
    def _nav_items(self):
        if self.role == ROLE_ADMIN:
            return [("users", "用户管理"), ("loginlogs", "登录日志")]
        if self.role == ROLE_OPERATOR:
            return [("devices", "设备管理")]
        return [("oplogs", "操作日志")]

    def _default_view(self):
        if self.role == ROLE_OPERATOR:
            return "devices"
        if self.role == ROLE_ADMIN:
            return "users"
        return "oplogs"

    def _build_nav(self):
        self.nav_frame = tk.Frame(self.root, bg=self.C_BG)
        self.nav_frame.pack(side=tk.TOP, fill=tk.X, padx=12, pady=(8, 0))
        self._nav_buttons = {}
        for key, label in self._nav_items():
            self._nav_buttons[key] = ttk.Button(self.nav_frame, text=label,
                                                command=lambda k=key: self.show_view(k))
        self.nav_frame.lift()

    def _update_nav_styles(self):
        for key, btn in self._nav_buttons.items():
            btn.configure(style="NavSel.TButton" if key == self._current_view else "Nav.TButton")
            btn.pack(side=tk.LEFT, padx=(0, 6))

    # ---------- 视图切换 ----------
    def show_view(self, name):
        self._current_view = name
        clear_children(self.content)
        for btn in self._nav_buttons.values():
            btn.pack_forget()
        self._update_nav_styles()
        if name == "devices":
            self._build_device_view()
        elif name == "users":
            self._build_users_view()
        elif name == "loginlogs":
            self._build_login_logs_view()
        elif name == "oplogs":
            self._build_op_logs_view()

    def _make_tree(self, columns, widths, center_cols=()):
        wrap = tk.Frame(self.content, bg=self.C_CARD, highlightbackground=self.C_BORDER,
                        highlightthickness=1)
        wrap.pack(fill=tk.BOTH, expand=True, pady=(6, 2))
        tree = ttk.Treeview(wrap, columns=columns, show="headings", selectmode="browse")
        yscroll = ttk.Scrollbar(wrap, orient=tk.VERTICAL, command=tree.yview,
                                style="Vertical.TScrollbar")
        tree.configure(yscrollcommand=yscroll.set)
        yscroll.pack(side=tk.RIGHT, fill=tk.Y)
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        for col in columns:
            tree.heading(col, text=col)
            anchor = tk.CENTER if col in center_cols else tk.W
            tree.column(col, width=widths.get(col, 110), anchor=anchor, stretch=True)
        tree.tag_configure("odd", background=self.C_ZEBRA)
        tree.tag_configure("even", background=self.C_CARD)
        return tree

    def _build_toolbar_row(self):
        bar = tk.Frame(self.content, bg=self.C_BG)
        bar.pack(fill=tk.X, pady=(2, 0))
        return bar

    # ==========================================
    # 操作员：设备管理
    # ==========================================
    def _build_device_view(self):
        bar = self._build_toolbar_row()
        ttk.Button(bar, text="＋ 添加设备", style="Tool.TButton",
                   command=self._device_dialog).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(bar, text="编辑选中", style="Tool.TButton",
                   command=self._edit_selected_device).pack(side=tk.LEFT, padx=6)
        ttk.Button(bar, text="× 删除选中", style="Danger.TButton",
                   command=self._delete_selected_device).pack(side=tk.LEFT, padx=6)
        ttk.Button(bar, text="导出Excel", style="Ghost.TButton",
                   command=self._export_devices).pack(side=tk.LEFT, padx=6)
        ttk.Button(bar, text="导入Excel", style="Ghost.TButton",
                   command=self._import_devices).pack(side=tk.LEFT, padx=6)

        right = tk.Frame(bar, bg=self.C_BG)
        right.pack(side=tk.RIGHT)
        ttk.Checkbutton(right, text="显示明文密码", variable=self.show_pwd,
                        command=self._refresh_devices).pack(side=tk.LEFT, padx=(0, 10))
        tk.Label(right, text="搜索：", bg=self.C_BG, fg=self.C_TEXT,
                 font=(self.font_family, 10)).pack(side=tk.LEFT)
        self.dev_search = tk.StringVar()
        self.dev_search.trace("w", lambda *a: self._refresh_devices())
        ttk.Entry(right, textvariable=self.dev_search, style="Search.TEntry",
                  width=24).pack(side=tk.LEFT)

        self.dev_tree = self._make_tree(
            ("ID", "机房名称", "机柜位置", "设备名称", "类型", "IP地址", "端口", "账号", "密码", "备注"),
            {"ID": 55, "机房名称": 105, "机柜位置": 105, "设备名称": 140, "类型": 85,
             "IP地址": 128, "端口": 60, "账号": 95, "密码": 105, "备注": 170},
            ("ID", "类型", "端口"))
        self.dev_tree.bind("<Double-1>", lambda e: self._edit_selected_device())
        self._device_menu = tk.Menu(self.root, tearoff=0, bg=self.C_CARD, fg=self.C_TEXT,
                                    activebackground=self.C_SELECT, activeforeground=self.C_TEXT,
                                    bd=1, relief="solid", font=(self.font_family, 10))
        self._device_menu.add_command(label="编辑设备", command=self._edit_selected_device)
        self._device_menu.add_command(label="删除设备", command=self._delete_selected_device)
        self._device_menu.add_separator()
        self._device_menu.add_command(label="显示/隐藏明文密码", command=self._toggle_show_pwd)
        self.dev_tree.bind("<Button-3>", self._show_dev_menu)

        self._refresh_devices()

    def _show_dev_menu(self, event):
        try:
            item = self.dev_tree.identify_row(event.y)
            if item:
                self.dev_tree.selection_set(item)
                self._device_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self._device_menu.grab_release()

    def _toggle_show_pwd(self):
        self.show_pwd.set(not self.show_pwd.get())
        self._refresh_devices()

    def _refresh_devices(self):
        for item in self.dev_tree.get_children():
            self.dev_tree.delete(item)
        search = self.dev_search.get().strip().lower()
        show_plain = self.show_pwd.get()
        rows = fetch_all_devices()
        self.total_count = len(rows)
        idx = 0
        for row in rows:
            did, room, rack, name, type_, ip, port, account, pwd_enc, notes = row
            pwd_plain = decrypt_text(self.key, pwd_enc)
            if search:
                hay = " ".join(str(x) for x in (room, rack, name, type_, ip, port, account,
                                                pwd_plain, notes)).lower()
                if search not in hay:
                    continue
            display_pwd = pwd_plain if show_plain else ("******" if pwd_enc else "")
            self.dev_tree.insert("", tk.END, iid=str(did), values=(
                did, room, rack, name, type_, ip, port, account, display_pwd, notes),
                tags=("odd" if idx % 2 else "even",))
            idx += 1
        self.filtered_count = idx
        self._update_status("设备共 %d 条" % self.total_count +
                            (" ｜ 筛选 %d 条" % idx if search else ""))

    def _selected_device_id(self):
        sel = self.dev_tree.selection()
        if not sel:
            messagebox.showwarning("提示", "请先选中一行数据", parent=self.root)
            return None
        return int(self.dev_tree.item(sel[0])["values"][0])

    def _device_dialog(self, device_id=None):
        dialog = tk.Toplevel(self.root)
        dialog.title("添加设备" if device_id is None else "编辑设备")
        dialog.configure(bg=self.C_BG)
        dialog.transient(self.root)
        dialog.grab_set()
        dialog.resizable(False, False)

        fields = [("机房名称", True), ("机柜位置", True), ("设备名称", True), ("类型", True),
                  ("IP地址", False), ("端口", False), ("账号", False), ("密码", False), ("备注", False)]
        entries = {}
        type_var = tk.StringVar()

        data = get_device_by_id(device_id) if device_id is not None else None

        # 底部按钮区先打包、锚定在底部，保证可见
        btns = tk.Frame(dialog, bg=self.C_BG)
        btns.pack(side=tk.BOTTOM, fill=tk.X, padx=20, pady=16)

        body = tk.Frame(dialog, bg=self.C_BG)
        body.pack(fill=tk.BOTH, expand=True, padx=20, pady=(16, 0))
        for i, (label, required) in enumerate(fields):
            tk.Label(body, text=label + (" *" if required else ""), bg=self.C_BG, fg=self.C_TEXT,
                     font=(self.font_family, 10)).grid(row=i, column=0, sticky=tk.W, padx=(0, 12), pady=5)
            if label == "类型":
                combo = ttk.Combobox(body, textvariable=type_var, values=DEVICE_TYPES,
                                     state="readonly", width=28)
                type_var.set(data[3] if data else DEVICE_TYPES[0])
                combo.grid(row=i, column=1, sticky=tk.W, pady=6, ipady=3)
                entries[label] = combo
            else:
                entry = ttk.Entry(body, width=30, font=(self.font_family, 11))
                if data:
                    vi = [f[0] for f in fields].index(label)
                    val = data[vi]
                    if label == "密码":
                        val = decrypt_text(self.key, val)
                    entry.insert(0, val if val else "")
                entry.grid(row=i, column=1, sticky=tk.W, pady=6, ipady=4)
                entries[label] = entry
        tk.Label(body, text="* 为必填字段；密码加密存储，列表默认脱敏", bg=self.C_BG, fg=self.C_MUTED,
                 font=(self.font_family, 9)).grid(row=len(fields), column=0, columnspan=2,
                                                  sticky=tk.W, pady=(6, 0))

        def save():
            d = {}
            for label, entry in entries.items():
                d[label] = entry.get().strip() if not isinstance(entry, ttk.Combobox) else entry.get()
            for label, required in fields:
                if required and not d[label]:
                    messagebox.showerror("错误", "字段“%s”不能为空！" % label, parent=dialog)
                    return
            pwd_enc = encrypt_text(self.key, d["密码"])
            if device_id is None:
                add_device(d["机房名称"], d["机柜位置"], d["设备名称"], d["类型"], d["IP地址"],
                           d["端口"], d["账号"], pwd_enc, d["备注"])
                op_log(self.user["username"], self.role_cn, "添加设备", d["设备名称"],
                       "%s / %s" % (d["机房名称"], d["机柜位置"]))
            else:
                update_device(device_id, d["机房名称"], d["机柜位置"], d["设备名称"], d["类型"],
                              d["IP地址"], d["端口"], d["账号"], pwd_enc, d["备注"])
                op_log(self.user["username"], self.role_cn, "编辑设备", d["设备名称"],
                       "%s / %s" % (d["机房名称"], d["机柜位置"]))
            self._refresh_devices()
            dialog.destroy()

        tk.Button(btns, text="添加设备" if device_id is None else "保 存", command=save,
                  bg="#1E88E5", fg="#FFFFFF", activebackground="#1565C0",
                  activeforeground="#FFFFFF", relief="flat", padx=20, pady=6,
                  font=(self.font_family, 11, "bold")).pack(side=tk.RIGHT, padx=(8, 0))
        tk.Button(btns, text="取 消", command=dialog.destroy,
                  bg="#FFFFFF", fg="#1E293B", activebackground="#E2E8F0",
                  relief="solid", bd=1, padx=20, pady=6,
                  font=(self.font_family, 10)).pack(side=tk.RIGHT)

        _fit_center(dialog, min_h=470)

    def _edit_selected_device(self):
        did = self._selected_device_id()
        if did is not None:
            self._device_dialog(did)

    def _delete_selected_device(self):
        did = self._selected_device_id()
        if did is None:
            return
        dev = get_device_by_id(did)
        name = dev[2] if dev else "该设备"
        if messagebox.askyesno("确认删除", "确定要删除设备“%s”吗？\n此操作不可恢复。" % name,
                               icon="warning", parent=self.root):
            delete_device(did)
            op_log(self.user["username"], self.role_cn, "删除设备", name, "")
            self._refresh_devices()
            messagebox.showinfo("成功", "设备已删除", parent=self.root)

    def _export_devices(self):
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("缺少依赖", "未安装 openpyxl，无法导出 Excel。\n\n请执行： pip3 install openpyxl",
                                 parent=self.root)
            return
        data = fetch_all_devices()
        if not data:
            messagebox.showwarning("提示", "没有数据可导出", parent=self.root)
            return
        include_plain = messagebox.askyesno(
            "导出选项",
            "是否导出明文密码？\n\n  [是] 明文（交接/备份）\n  [否] 脱敏（******）\n\n对外提供建议选 [否]。",
            icon="question", default="no", parent=self.root)
        path = filedialog.asksaveasfilename(parent=self.root, defaultextension=".xlsx",
                                            filetypes=[("Excel文件", "*.xlsx")], title="导出Excel")
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "设备密码"
        ws.append(["ID", "机房名称", "机柜位置", "设备名称", "类型", "IP地址", "端口", "账号", "密码", "备注"])
        for row in data:
            r = list(row)
            r[8] = decrypt_text(self.key, r[8]) if include_plain else ("******" if r[8] else "")
            ws.append(r)
        wb.save(path)
        op_log(self.user["username"], self.role_cn, "导出Excel", "", "%d 条记录" % len(data))
        messagebox.showinfo("成功", "数据已导出到：\n%s" % path, parent=self.root)

    def _import_devices(self):
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("缺少依赖", "未安装 openpyxl，无法导入 Excel。\n\n请执行： pip3 install openpyxl",
                                 parent=self.root)
            return
        path = filedialog.askopenfilename(parent=self.root, filetypes=[("Excel文件", "*.xlsx")],
                                          title="选择Excel文件")
        if not path:
            return
        try:
            wb = load_workbook(path)
            ws = wb.active
        except Exception as e:
            messagebox.showerror("错误", "无法读取Excel文件：%s" % e, parent=self.root)
            return
        headers = [c.value for c in ws[1]]
        expected = ["机房名称", "机柜位置", "设备名称", "类型", "IP地址", "端口", "账号", "密码", "备注"]
        if headers[:4] != expected[:4]:
            messagebox.showerror("格式错误",
                                 "Excel格式不正确，请使用模板（表头：机房名称、机柜位置、设备名称、类型、IP地址、端口、账号、密码、备注）",
                                 parent=self.root)
            return
        count = 0
        errors = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or all(c is None for c in row):
                continue
            row = list(row) + [""] * (9 - len(row))
            data_row = row[1:] if headers[0] == "ID" else row[:9]
            room, rack, name, type_, ip, port, account, password, notes = data_row
            if not room or not rack or not name or not type_:
                errors.append("第%d行：必填字段为空，已跳过" % row_idx)
                continue
            if type_ not in DEVICE_TYPES:
                errors.append("第%d行：类型无效，已跳过" % row_idx)
                continue
            add_device(str(room), str(rack), str(name), str(type_), str(ip or ""), str(port or ""),
                       str(account or ""), encrypt_text(self.key, str(password or "")), str(notes or ""))
            count += 1
        self._refresh_devices()
        op_log(self.user["username"], self.role_cn, "导入Excel", "", "%d 条记录" % count)
        msg = "成功导入 %d 条记录" % count
        if errors:
            msg += "\n\n警告：" + "\n".join(errors[:5])
        messagebox.showinfo("导入完成", msg, parent=self.root)

    # ==========================================
    # 管理员：用户管理
    # ==========================================
    def _build_users_view(self):
        bar = self._build_toolbar_row()
        ttk.Button(bar, text="＋ 添加用户", style="Tool.TButton",
                   command=self._user_dialog).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(bar, text="编辑用户", style="Tool.TButton",
                   command=self._edit_selected_user).pack(side=tk.LEFT, padx=6)
        ttk.Button(bar, text="重置密码", style="Ghost.TButton",
                   command=self._reset_selected_pwd).pack(side=tk.LEFT, padx=6)
        ttk.Button(bar, text="× 删除用户", style="Danger.TButton",
                   command=self._delete_selected_user).pack(side=tk.LEFT, padx=6)

        right = tk.Frame(bar, bg=self.C_BG)
        right.pack(side=tk.RIGHT)
        tk.Label(right, text="搜索：", bg=self.C_BG, fg=self.C_TEXT,
                 font=(self.font_family, 10)).pack(side=tk.LEFT)
        self.user_search = tk.StringVar()
        self.user_search.trace("w", lambda *a: self._refresh_users())
        ttk.Entry(right, textvariable=self.user_search, style="Search.TEntry",
                  width=24).pack(side=tk.LEFT)

        self.user_tree = self._make_tree(
            ("用户名", "角色", "状态", "创建时间", "创建人", "最近登录"),
            {"用户名": 150, "角色": 90, "状态": 80, "创建时间": 160, "创建人": 120, "最近登录": 160},
            ("角色", "状态"))
        self.user_tree.bind("<Double-1>", lambda e: self._edit_selected_user())
        self._refresh_users()

    def _refresh_users(self):
        for item in self.user_tree.get_children():
            self.user_tree.delete(item)
        search = self.user_search.get().strip().lower()
        rows = list_users()
        self.total_count = len(rows)
        idx = 0
        for uid, uname, role, enabled, created_at, created_by, last_login in rows:
            if search and search not in uname.lower() and search not in ROLE_CN.get(role, "").lower():
                continue
            status = "启用" if enabled else "禁用"
            vals = (uname, ROLE_CN.get(role, role), status, created_at or "", created_by or "", last_login or "")
            self.user_tree.insert("", tk.END, iid=str(uid), values=vals,
                                  tags=("odd" if idx % 2 else "even",))
            idx += 1
        self.filtered_count = idx
        self._update_status("用户共 %d 个" % self.total_count +
                            (" ｜ 筛选 %d 个" % idx if search else ""))

    def _selected_user_row(self):
        sel = self.user_tree.selection()
        if not sel:
            messagebox.showwarning("提示", "请先选中一行数据", parent=self.root)
            return None
        uid = int(sel[0])
        for r in list_users():
            if r[0] == uid:
                return r
        return None

    def _user_dialog(self, user_id=None):
        editing = user_id is not None
        target = None
        for r in list_users():
            if r[0] == user_id:
                target = r
                break
        if editing and target is None:
            messagebox.showerror("错误", "未找到该用户记录", parent=self.root)
            return

        dialog = tk.Toplevel(self.root)
        dialog.title("添加用户" if not editing else "编辑用户")
        dialog.configure(bg=self.C_BG)
        dialog.transient(self.root)
        dialog.grab_set()
        dialog.resizable(False, False)

        # 固定尺寸 + 三重锁定（geometry/minsize/maxsize），彻底杜绝被压缩裁掉按钮
        dlg_w, dlg_h = 480, (540 if not editing else 440)
        sw, sh = dialog.winfo_screenwidth(), dialog.winfo_screenheight()
        x = (sw - dlg_w) // 2
        y = max((sh - dlg_h) // 3, 0)
        dialog.geometry("%dx%d+%d+%d" % (dlg_w, dlg_h, x, y))
        dialog.minsize(dlg_w, dlg_h)
        dialog.maxsize(dlg_w, dlg_h)

        # 底部按钮区先打包、锚定在底部
        btns = tk.Frame(dialog, bg=self.C_BG)
        btns.pack(side=tk.BOTTOM, fill=tk.X, padx=24, pady=14)

        # 顶部标题（带版本标记，便于确认是否为新版）
        tk.Label(dialog, text=("添加新用户" if not editing else "编辑用户") + "  [v3.1]",
                 bg=self.C_BG, fg=self.C_PRIMARY,
                 font=(self.font_family, 13, "bold")).pack(anchor=tk.W, padx=24, pady=(14, 0))

        # 表单区：grid 布局，每行一个字段，位置绝对确定
        body = tk.Frame(dialog, bg=self.C_BG)
        body.pack(fill=tk.BOTH, expand=True, padx=24, pady=(8, 0))
        body.grid_columnconfigure(0, weight=1)
        r = 0

        tk.Label(body, text="用户名 *", bg=self.C_BG, fg=self.C_TEXT,
                 font=(self.font_family, 11)).grid(row=r, column=0, sticky=tk.W, pady=(6, 2))
        r += 1
        uname_var = tk.StringVar(value=target[1] if target else "")
        uname_entry = tk.Entry(body, textvariable=uname_var, font=(self.font_family, 12),
                               relief="solid", bd=1, state="readonly" if editing else "normal")
        uname_entry.grid(row=r, column=0, sticky=tk.EW, pady=(0, 6))
        r += 1

        tk.Label(body, text="角色 *", bg=self.C_BG, fg=self.C_TEXT,
                 font=(self.font_family, 11)).grid(row=r, column=0, sticky=tk.W, pady=(6, 2))
        r += 1
        role_var = tk.StringVar(value=ROLE_CN[target[2]] if target else ROLE_CN[ROLE_OPERATOR])
        role_combo = ttk.Combobox(body, textvariable=role_var,
                                  values=[ROLE_CN[r] for r in ALL_ROLES], state="readonly",
                                  font=(self.font_family, 12))
        role_combo.grid(row=r, column=0, sticky=tk.EW, pady=(0, 6))
        r += 1

        if editing:
            status_var = tk.StringVar(value="启用" if target[3] else "禁用")
            tk.Label(body, text="状态", bg=self.C_BG, fg=self.C_TEXT,
                     font=(self.font_family, 11)).grid(row=r, column=0, sticky=tk.W, pady=(6, 2))
            r += 1
            ttk.Combobox(body, textvariable=status_var, values=["启用", "禁用"],
                         state="readonly", font=(self.font_family, 12)).grid(
                             row=r, column=0, sticky=tk.EW, pady=(0, 6))
            r += 1

        # 密码（添加时才显示）
        if not editing:
            tk.Label(body, text="初始密码 *（至少%d位，含字母和数字）" % MIN_PASSWORD_LEN,
                     bg=self.C_BG, fg=self.C_TEXT, font=(self.font_family, 11)).grid(
                         row=r, column=0, sticky=tk.W, pady=(6, 2))
            r += 1
            pwd_var = tk.StringVar()
            tk.Entry(body, textvariable=pwd_var, show="*", font=(self.font_family, 12),
                     relief="solid", bd=1).grid(row=r, column=0, sticky=tk.EW, pady=(0, 6))
            r += 1
            tk.Label(body, text="确认密码 *", bg=self.C_BG, fg=self.C_TEXT,
                     font=(self.font_family, 11)).grid(row=r, column=0, sticky=tk.W, pady=(6, 2))
            r += 1
            cfm_var = tk.StringVar()
            tk.Entry(body, textvariable=cfm_var, show="*", font=(self.font_family, 12),
                     relief="solid", bd=1).grid(row=r, column=0, sticky=tk.EW, pady=(0, 6))
            r += 1

        tk.Label(body, text="角色说明：管理员管用户、操作员管设备、审计员看操作日志",
                 bg=self.C_BG, fg=self.C_MUTED, font=(self.font_family, 9)).grid(
                     row=r, column=0, sticky=tk.W, pady=(8, 0))

        def save():
            uname = uname_var.get().strip()
            role_cn = role_var.get()
            role = [k for k, v in ROLE_CN.items() if v == role_cn][0]
            if not uname:
                messagebox.showerror("错误", "用户名不能为空！", parent=dialog)
                return
            if not editing:
                if username_exists(uname):
                    messagebox.showerror("错误", "用户名“%s”已存在！" % uname, parent=dialog)
                    return
                if not _validate_password_dialog(dialog, pwd_var.get(), cfm_var.get()):
                    return
                create_user(uname, pwd_var.get(), role, self.user["username"])
                op_log(self.user["username"], self.role_cn, "创建用户", uname, "角色：" + role_cn)
            else:
                if target[2] == ROLE_ADMIN and target[3] == 1:
                    if role != ROLE_ADMIN:
                        if enabled_admin_count() <= 1:
                            messagebox.showerror("禁止操作", "系统必须保留至少一个启用的管理员！", parent=dialog)
                            return
                    if status_var.get() == "禁用":
                        if enabled_admin_count() <= 1:
                            messagebox.showerror("禁止操作", "系统必须保留至少一个启用的管理员！", parent=dialog)
                            return
                if role != target[2]:
                    update_user_role(user_id, role)
                    op_log(self.user["username"], self.role_cn, "修改用户角色", uname,
                           "%s → %s" % (ROLE_CN[target[2]], role_cn))
                new_status = status_var.get() == "启用"
                if new_status != bool(target[3]):
                    set_user_enabled(user_id, new_status)
                    op_log(self.user["username"], self.role_cn,
                           "启用用户" if new_status else "禁用用户", uname, "")
            self._refresh_users()
            dialog.destroy()

        tk.Button(btns, text="添加用户" if not editing else "保 存", command=save,
                  bg="#1E88E5", fg="#FFFFFF", activebackground="#1565C0",
                  activeforeground="#FFFFFF", relief="flat", padx=20, pady=6,
                  font=(self.font_family, 11, "bold")).pack(side=tk.RIGHT, padx=(8, 0))
        tk.Button(btns, text="取 消", command=dialog.destroy,
                  bg="#FFFFFF", fg="#1E293B", activebackground="#E2E8F0",
                  relief="solid", bd=1, padx=20, pady=6,
                  font=(self.font_family, 10)).pack(side=tk.RIGHT)

    def _edit_selected_user(self):
        row = self._selected_user_row()
        if row:
            self._user_dialog(row[0])

    def _reset_selected_pwd(self):
        row = self._selected_user_row()
        if not row:
            return
        uid, uname = row[0], row[1]
        dialog = tk.Toplevel(self.root)
        dialog.title("重置密码 - %s" % uname)
        dialog.configure(bg=self.C_BG)
        dialog.transient(self.root)
        dialog.grab_set()
        dialog.resizable(False, False)

        # 底部按钮区先打包、锚定在底部，保证可见
        btns = tk.Frame(dialog, bg=self.C_BG)
        btns.pack(side=tk.BOTTOM, fill=tk.X, padx=22, pady=12)

        tk.Label(dialog, text="为用户“%s”设置新密码" % uname, bg=self.C_BG, fg=self.C_TEXT,
                 font=(self.font_family, 11, "bold")).pack(anchor=tk.W, padx=22, pady=(16, 4))
        tk.Label(dialog, text="密码至少 %d 位，且需同时包含字母和数字" % MIN_PASSWORD_LEN,
                 bg=self.C_BG, fg=self.C_MUTED, font=(self.font_family, 9)).pack(anchor=tk.W, padx=22)

        body = tk.Frame(dialog, bg=self.C_BG)
        body.pack(fill=tk.BOTH, expand=True, padx=22, pady=8)
        tk.Label(body, text="新密码 *", bg=self.C_BG, fg=self.C_TEXT,
                 font=(self.font_family, 10)).pack(anchor=tk.W, pady=(4, 2))
        pwd_var = tk.StringVar()
        tk.Entry(body, textvariable=pwd_var, show="*", width=30, font=(self.font_family, 11),
                 relief="solid", bd=1).pack(fill=tk.X, pady=(0, 6))
        tk.Label(body, text="确认新密码 *", bg=self.C_BG, fg=self.C_TEXT,
                 font=(self.font_family, 10)).pack(anchor=tk.W, pady=(4, 2))
        cfm_var = tk.StringVar()
        tk.Entry(body, textvariable=cfm_var, show="*", width=30, font=(self.font_family, 11),
                 relief="solid", bd=1).pack(fill=tk.X, pady=(0, 6))

        def save():
            if not _validate_password_dialog(dialog, pwd_var.get(), cfm_var.get()):
                return
            set_user_password(uid, pwd_var.get())
            op_log(self.user["username"], self.role_cn, "重置密码", uname, "")
            dialog.destroy()
            messagebox.showinfo("成功", "用户“%s”的密码已重置。" % uname, parent=self.root)

        tk.Button(btns, text="确 定", command=save,
                  bg="#1E88E5", fg="#FFFFFF", activebackground="#1565C0",
                  activeforeground="#FFFFFF", relief="flat", padx=18, pady=6,
                  font=(self.font_family, 11, "bold")).pack(side=tk.RIGHT, padx=(8, 0))
        tk.Button(btns, text="取 消", command=dialog.destroy,
                  bg="#FFFFFF", fg="#1E293B", activebackground="#E2E8F0",
                  relief="solid", bd=1, padx=18, pady=6,
                  font=(self.font_family, 10)).pack(side=tk.RIGHT)

        _fit_center(dialog, min_h=300)

    def _delete_selected_user(self):
        row = self._selected_user_row()
        if not row:
            return
        uid, uname, role = row[0], row[1], row[2]
        if uname == self.user["username"]:
            messagebox.showerror("禁止操作", "不能删除当前登录的账号！", parent=self.root)
            return
        if role == ROLE_ADMIN and row[3] == 1 and enabled_admin_count() <= 1:
            messagebox.showerror("禁止操作", "系统必须保留至少一个启用的管理员，不能删除！", parent=self.root)
            return
        if messagebox.askyesno("确认删除", "确定要删除用户“%s”（%s）吗？\n此操作不可恢复。" % (uname, ROLE_CN.get(role, role)),
                               icon="warning", parent=self.root):
            delete_user(uid)
            op_log(self.user["username"], self.role_cn, "删除用户", uname, "")
            self._refresh_users()
            messagebox.showinfo("成功", "用户已删除", parent=self.root)

    # ==========================================
    # 管理员：登录日志（只能看登录日志）
    # ==========================================
    def _build_login_logs_view(self):
        bar = self._build_toolbar_row()
        ttk.Button(bar, text="刷新", style="Ghost.TButton",
                   command=self._refresh_login_logs).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(bar, text="清空登录日志", style="Danger.TButton",
                   command=self._clear_login_logs).pack(side=tk.LEFT, padx=6)

        right = tk.Frame(bar, bg=self.C_BG)
        right.pack(side=tk.RIGHT)
        tk.Label(right, text="搜索：", bg=self.C_BG, fg=self.C_TEXT,
                 font=(self.font_family, 10)).pack(side=tk.LEFT)
        self.ll_search = tk.StringVar()
        self.ll_search.trace("w", lambda *a: self._refresh_login_logs())
        ttk.Entry(right, textvariable=self.ll_search, style="Search.TEntry",
                  width=24).pack(side=tk.LEFT)

        self.ll_tree = self._make_tree(
            ("时间", "用户", "角色", "结果", "详情"),
            {"时间": 160, "用户": 140, "角色": 80, "结果": 110, "详情": 220},
            ("角色",))
        self._refresh_login_logs()

    def _refresh_login_logs(self):
        for item in self.ll_tree.get_children():
            self.ll_tree.delete(item)
        rows = list_login_logs(self.ll_search.get().strip())
        self.total_count = len(rows)
        idx = 0
        for ts, uname, role, result, detail in rows:
            self.ll_tree.insert("", tk.END, iid=str(idx), values=(
                ts, uname, role, result, detail), tags=("odd" if idx % 2 else "even",))
            idx += 1
        self.filtered_count = idx
        self._update_status("登录日志共 %d 条" % self.total_count)

    def _clear_login_logs(self):
        if messagebox.askyesno("清空日志", "确定要清空全部登录日志吗？\n此操作不可恢复。",
                               icon="warning", parent=self.root):
            clear_login_logs()
            op_log(self.user["username"], self.role_cn, "清空登录日志", "", "")
            self._refresh_login_logs()
            messagebox.showinfo("成功", "登录日志已清空", parent=self.root)

    # ==========================================
    # 审计员：操作日志（只读 + 可导出）
    # ==========================================
    def _build_op_logs_view(self):
        bar = self._build_toolbar_row()
        ttk.Button(bar, text="刷新", style="Ghost.TButton",
                   command=self._refresh_op_logs).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(bar, text="导出CSV", style="Tool.TButton",
                   command=self._export_op_csv).pack(side=tk.LEFT, padx=6)
        export_btn = ttk.Button(bar, text="导出Excel", style="Tool.TButton",
                                command=self._export_op_excel)
        export_btn.pack(side=tk.LEFT, padx=6)
        if not OPENPYXL_AVAILABLE:
            export_btn.state(["disabled"])
            tk.Label(bar, text="（未安装openpyxl）", bg=self.C_BG, fg=self.C_MUTED,
                     font=(self.font_family, 9)).pack(side=tk.LEFT, padx=(0, 8))

        right = tk.Frame(bar, bg=self.C_BG)
        right.pack(side=tk.RIGHT)
        tk.Label(right, text="搜索：", bg=self.C_BG, fg=self.C_TEXT,
                 font=(self.font_family, 10)).pack(side=tk.LEFT)
        self.op_search = tk.StringVar()
        self.op_search.trace("w", lambda *a: self._refresh_op_logs())
        ttk.Entry(right, textvariable=self.op_search, style="Search.TEntry",
                  width=24).pack(side=tk.LEFT)

        self.op_tree = self._make_tree(
            ("时间", "用户", "角色", "操作", "对象", "详情"),
            {"时间": 160, "用户": 120, "角色": 80, "操作": 110, "对象": 130, "详情": 200},
            ("角色",))
        self._refresh_op_logs()

    def _refresh_op_logs(self):
        for item in self.op_tree.get_children():
            self.op_tree.delete(item)
        rows = list_operation_logs(self.op_search.get().strip())
        self.total_count = len(rows)
        idx = 0
        for ts, uname, role, action, target, detail in rows:
            self.op_tree.insert("", tk.END, iid=str(idx), values=(
                ts, uname, role, action, target, detail), tags=("odd" if idx % 2 else "even",))
            idx += 1
        self.filtered_count = idx
        self._update_status("操作日志共 %d 条" % self.total_count)

    def _export_op_csv(self):
        path = filedialog.asksaveasfilename(parent=self.root, defaultextension=".csv",
                                            filetypes=[("CSV文件", "*.csv")], title="导出操作日志(CSV)")
        if not path:
            return
        try:
            n = export_operation_logs_to_csv(path)
            messagebox.showinfo("成功", "已导出 %d 条操作日志到：\n%s" % (n, path), parent=self.root)
        except Exception as e:
            messagebox.showerror("错误", "导出失败：%s" % e, parent=self.root)

    def _export_op_excel(self):
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("缺少依赖", "未安装 openpyxl，无法导出 Excel。\n\n请执行： pip3 install openpyxl",
                                 parent=self.root)
            return
        path = filedialog.asksaveasfilename(parent=self.root, defaultextension=".xlsx",
                                            filetypes=[("Excel文件", "*.xlsx")], title="导出操作日志(Excel)")
        if not path:
            return
        try:
            n = export_operation_logs_to_excel(path)
            messagebox.showinfo("成功", "已导出 %d 条操作日志到：\n%s" % (n, path), parent=self.root)
        except Exception as e:
            messagebox.showerror("错误", "导出失败：%s" % e, parent=self.root)

    # ---------- 状态栏 ----------
    def _build_statusbar(self):
        sb = tk.Frame(self.root, bg=self.C_HEADER, height=30)
        sb.pack(side=tk.BOTTOM, fill=tk.X)
        sb.pack_propagate(False)
        self.stat_label = tk.Label(sb, text="", bg=self.C_HEADER, fg="#CBD5E1",
                                   font=(self.font_family, 9))
        self.stat_label.pack(side=tk.LEFT, padx=14, pady=6)
        tk.Label(sb, text="当前角色：%s" % self.role_cn, bg=self.C_HEADER, fg=ROLE_COLORS[self.role],
                 font=(self.font_family, 9)).pack(side=tk.RIGHT, padx=14, pady=6)

    def _update_status(self, text):
        self.stat_label.config(text=text)


# ============================================================
# 八、主程序入口
# ============================================================
def main():
    init_db()
    key = get_or_create_key()
    migrated = migrate_plaintext_passwords(key)
    if migrated:
        messagebox.showinfo("数据安全升级", "已将 %d 条旧明文设备密码升级为加密存储。" % migrated)

    root = tk.Tk()
    root.withdraw()

    # 首次运行：强制创建管理员账号
    if count_users() == 0:
        if not initial_admin_dialog(root):
            root.destroy()
            sys.exit(0)

    # 登录 → 使用 → （退出登录则回到登录界面）
    while True:
        clear_children(root)
        user = login_dialog(root)
        if user is None:
            root.destroy()
            sys.exit(0)
        root.deiconify()
        clear_children(root)
        app = DeviceManagerApp(root, user)
        root.mainloop()
        if not getattr(app, "_logout_requested", False):
            break
    root.destroy()


if __name__ == "__main__":
    main()
